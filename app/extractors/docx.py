from pathlib import Path
from typing import List, Dict, Any
import docx


class DOCXExtractor:
    @staticmethod
    def extract(file_path: str) -> List[Dict[str, Any]]:
        """
        Extract text from DOCX document paragraphs and tables.
        Returns a list of dicts: [{'text': str, 'page_no': None}]
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        extracted_lines = []
        try:
            doc = docx.Document(str(path))
            for p in doc.paragraphs:
                clean_text = p.text.strip()
                if clean_text:
                    extracted_lines.append(clean_text)

            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        extracted_lines.append(" | ".join(row_texts))
        except Exception as e:
            raise RuntimeError(f"Failed to extract DOCX ({path.name}): {str(e)}") from e

        full_text = "\n\n".join(extracted_lines).strip()
        if not full_text:
            return []

        return [{"text": full_text, "page_no": None}]
